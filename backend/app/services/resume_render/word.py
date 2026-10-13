'''
The Word file (KAN-207): one plain python-docx layout from the same view and
the same field text as the PDF (fields.py), on a Letter page. It is for job
applications that require Word, not a copy of Jake's look: one font, bold
headings, a tab stop for dates, built-in bullet paragraphs. Not stdlib, so
the Tectonic build script never imports it.
'''

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.services.resume_render import Identity
from app.services.resume_render.fields import (
    SECTION_LABELS,
    SKILLS_LABEL,
    TWO_ROW,
    details,
    heading,
    link_text,
    skill_lines,
)
from app.services.resume_view import ResumeView, ViewEntry

PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(0.6)
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN


def _paragraph(doc, *, space_after: float = 0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def _row(doc, left: str, right: str, *, bold: bool = False, italic: bool = False, tag: str = "") -> None:
    '''One heading row: text on the left, a right-aligned tab for the right.'''
    p = _paragraph(doc)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(left)
    run.bold, run.italic = bold, italic
    if tag:
        p.add_run(" | ")
        p.add_run(tag).italic = True
    if right:
        p.add_run("\t" + right).italic = italic


def _rule_below(paragraph) -> None:
    # A bottom border, the plain stand-in for Jake's \titlerule
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _section_heading(doc, label: str) -> None:
    p = _paragraph(doc, space_after=2)
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(label.upper())
    run.bold = True
    run.font.size = Pt(12)
    _rule_below(p)


def _entry(doc, entry: ViewEntry) -> None:
    h = heading(entry)
    if entry.section in TWO_ROW:
        _row(doc, h.title, h.right, bold=True)
        if h.subtitle or h.subright:
            _row(doc, h.subtitle, h.subright, italic=True)
    else:
        _row(doc, h.title, h.right, bold=True, tag=h.tag)
    for text in [line for _, line in details(entry)] + [b.text for b in entry.bullets]:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(4)


def render_docx(view: ResumeView, identity: Identity) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, MARGIN)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    name = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(identity.name)
    name.bold = True
    name.font.size = Pt(20)
    contact = [x for x in (identity.phone, identity.email) if x] + [link_text(u) for u in identity.links]
    if contact:
        _paragraph(doc, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(" | ".join(contact))

    for resume_section in view.sections:
        _section_heading(doc, SECTION_LABELS[resume_section.section])
        for entry in resume_section.entries:
            _entry(doc, entry)
    if view.skills:
        _section_heading(doc, SKILLS_LABEL)
        for label, skills in skill_lines(view.skills):
            p = _paragraph(doc)
            if label:
                p.add_run(f"{label}: ").bold = True
            p.add_run(", ".join(s.name for s in skills))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
