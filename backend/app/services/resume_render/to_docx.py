'''
Blocks as a Word document, with python-docx, for the application forms that
ask for one. Same layout as the PDF, built from plain paragraphs only: a
right tab stop puts dates on the right (no table, which some applicant
tracking systems read out of order), and Word's own small caps and bullet
list style stand in for the LaTeX ones.

A .docx can't carry its font the way a PDF does, so it names Times New Roman,
which every machine that opens one has, rather than CMU Serif.
'''
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.services.resume_render.blocks import (
    Block,
    Bullet,
    Contact,
    EntryHeader,
    Heading,
    Name,
    SkillLine,
)

FONT = "Times New Roman"
# The original's indents (to_pdf.py has the measurements)
INDENT = Inches(0.15)
RIGHT = Inches(7.5 * 0.97 + 0.15)  # of the 7.5in between the margins
BULLET_TEXT = Inches(0.47)
BULLET_HANG = Inches(0.1)


def _rule_under(paragraph) -> None:
    '''A bottom border on the paragraph: the heading's rule, as text, not a
    drawn line, so nothing extra sits in the document's flow.'''
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _tight(paragraph, before: float = 0, after: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), 1.0


def _two_sided(doc, left: str, right: str, italic: bool, size: float, note: str = "",
               right_size: float | None = None, before: float = 0) -> None:
    p = doc.add_paragraph()
    _tight(p, before=before)
    p.paragraph_format.left_indent = INDENT
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT - INDENT, WD_TAB_ALIGNMENT.RIGHT)
    head = p.add_run(left)
    head.bold, head.italic, head.font.size = not italic, italic, Pt(size)
    if note:
        p.add_run(" | ").font.size = Pt(size)
        extra = p.add_run(note)
        extra.italic, extra.font.size = True, Pt(size)
    if right:
        tail = p.add_run(f"\t{right}")
        tail.italic, tail.font.size = italic, Pt(right_size or size)


def write_docx(blocks: list[Block]) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.5)
        section.top_margin = section.bottom_margin = Inches(0.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(11)
    # East Asian text falls back to this too, or Word picks its own face
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for block in blocks:
        match block:
            case Name(text):
                p = doc.add_paragraph()
                _tight(p, after=2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Plain bold, as LaTeX prints it (to_pdf.py says why)
                run = p.add_run(text)
                run.bold, run.font.size = True, Pt(24.88)
            case Contact(items):
                p = doc.add_paragraph()
                _tight(p, after=4)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(" | ".join(text for text, _ in items)).font.size = Pt(10)
            case Heading(text):
                p = doc.add_paragraph()
                _tight(p, before=8, after=3)
                run = p.add_run(text)
                run.font.small_caps, run.font.size = True, Pt(12)
                _rule_under(p)
            case EntryHeader(left, right, sub_left, sub_right, note):
                if sub_left or sub_right:
                    _two_sided(doc, left, right, False, 11, before=4)
                    _two_sided(doc, sub_left, sub_right, True, 10)
                else:
                    # A project's line: name and tools small, date full size
                    _two_sided(doc, left, right, False, 10, note, right_size=11, before=4)
            case Bullet(text):
                p = doc.add_paragraph(style="List Bullet")
                _tight(p, after=1)
                p.paragraph_format.left_indent = BULLET_TEXT
                p.paragraph_format.first_line_indent = -BULLET_HANG
                p.add_run(text).font.size = Pt(10)
            case SkillLine(label, text):
                p = doc.add_paragraph()
                _tight(p)
                p.paragraph_format.left_indent = INDENT
                head = p.add_run(label)
                head.bold, head.font.size = True, Pt(10)
                p.add_run(f": {text}").font.size = Pt(10)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
