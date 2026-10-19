"""
Seam-1 tests for Generate (KAN-207): base resumes through the API on real
Postgres. Content is read back from the Word download, which needs no
Tectonic; the PDF render's own tests are in test_resume_render.py, and the
render route's tests below skip without Tectonic (backend/CLAUDE.md, Tests).
"""

import datetime
import uuid

import httpx
import pytest
from sqlalchemy import text

from app.deps import get_current_account
from app.main import app
from app.schemas.auth import AccountType, AuthenticatedAccount

pytestmark = pytest.mark.anyio


async def _seed_applicant(session) -> uuid.UUID:
    applicant_id = uuid.uuid4()
    await session.execute(text("INSERT INTO auth.users (id) VALUES (:id)"), {"id": applicant_id})
    await session.execute(
        text(
            "INSERT INTO applicant_profiles (id, email, full_name, phone_number, github_url)"
            " VALUES (:id, :email, 'Sam Seeker', '555-010-2030', 'https://github.com/samseeker')"
        ),
        {"id": applicant_id, "email": f"sam{applicant_id.hex[:6]}@example.com"},
    )
    await session.commit()
    return applicant_id


def _sign_in(applicant_id: uuid.UUID) -> None:
    app.dependency_overrides[get_current_account] = lambda: AuthenticatedAccount(
        id=applicant_id,
        email="sam@example.com",
        full_name="Sam Seeker",
        account_type=AccountType.APPLICANT,
        onboarding_completed=True,
    )


@pytest.fixture
async def student(database):
    """A signed-in applicant and an API client. Yields (client, applicant URL)."""
    async with database() as session:
        applicant_id = await _seed_applicant(session)
    _sign_in(applicant_id)
    transport = httpx.ASGITransport(app=app)
    try:
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            yield client, f"/applicants/{applicant_id}"
    finally:
        app.dependency_overrides.pop(get_current_account, None)


async def test_generate_makes_a_base_resume_the_list_reads_back(student):
    client, me = student
    res = await client.post(f"{me}/built-resumes", json={"name": "General SWE"})

    assert res.status_code == 201
    made = res.json()
    assert made["name"] == "General SWE"
    assert made["template"] == "jakes"
    assert made["job_posting_id"] is None
    assert made["is_default"] is True  # the student's only resume

    listed = (await client.get(f"{me}/built-resumes")).json()
    assert [(r["id"], r["name"]) for r in listed] == [(made["id"], "General SWE")]


async def _fill_bank(client, me) -> None:
    """Entries posted out of order, so the resume's order is the bank's."""
    bank = f"{me}/bank"
    for body in (
        {"section": "projects", "title": "WorkIt", "organization": "CSE 416", "start_month": "2026-01",
         "end_kind": "present", "bullets": ["Built a resume builder with FastAPI"]},
        {"section": "experience", "title": "Tutor", "organization": "Campus Center", "location": "Stony Brook, NY",
         "start_month": "2023-09", "end_month": "2024-05", "bullets": ["Tutored 40 students in calculus"]},
        {"section": "experience", "title": "Software Intern", "organization": "Acme", "location": "New York, NY",
         "start_month": "2025-06", "end_month": "2025-08",
         "bullets": ["Cut page load time by 30%", "Wrote the billing service's tests"]},
        {"section": "education", "title": "Stony Brook University", "location": "Stony Brook, NY",
         "degree": "Bachelor of Science", "field_of_study": "Computer Science", "gpa": 3.8, "gpa_scale": 4.0,
         "coursework": ["Data Structures", "Algorithms"], "start_month": "2023-08",
         "end_kind": "expected", "end_month": "2027-05"},
        {"section": "awards", "title": "Dean's List", "organization": "Stony Brook University", "award_date": "2024-12"},
    ):
        assert (await client.post(f"{bank}/entries", json=body)).status_code == 201
    for skill in ({"name": "Python", "group": "Languages"}, {"name": "Docker", "group": "Tools"}):
        assert (await client.post(f"{bank}/skills", json=skill)).status_code == 201


def _docx_lines(content: bytes) -> list[str]:
    from io import BytesIO

    from docx import Document

    return [p.text for p in Document(BytesIO(content)).paragraphs if p.text.strip()]


async def _generate(client, me, name="Base Resume") -> str:
    res = await client.post(f"{me}/built-resumes", json={"name": name})
    assert res.status_code == 201, res.text
    return res.json()["id"]


async def test_the_word_file_has_every_entry_in_section_order_newest_first(student):
    client, me = student
    await _fill_bank(client, me)
    resume_id = await _generate(client, me, "General SWE")

    res = await client.get(f"{me}/built-resumes/{resume_id}/docx")
    assert res.status_code == 200
    assert res.headers["content-disposition"] == 'attachment; filename="General SWE.docx"'
    assert _docx_lines(res.content) == [
        "Sam Seeker",
        f"555-010-2030 | sam{me.split('/')[2].replace('-', '')[:6]}@example.com | github.com/samseeker",
        "EDUCATION",
        "Stony Brook University\tStony Brook, NY",
        "Bachelor of Science in Computer Science\tAug. 2023 – Expected May 2027",
        "GPA: 3.8/4.0",
        "Relevant Coursework: Data Structures, Algorithms",
        "EXPERIENCE",
        "Software Intern\tJune 2025 – Aug. 2025",
        "Acme\tNew York, NY",
        "Cut page load time by 30%",
        "Wrote the billing service's tests",
        "Tutor\tSep. 2023 – May 2024",
        "Campus Center\tStony Brook, NY",
        "Tutored 40 students in calculus",
        "PROJECTS",
        "WorkIt | CSE 416\tJan. 2026 – Present",
        "Built a resume builder with FastAPI",
        "AWARDS",
        "Dean's List | Stony Brook University\tDec. 2024",
        "TECHNICAL SKILLS",
        "Languages: Python",
        "Tools: Docker",
    ]


async def test_the_word_file_is_letter_size(student):
    from io import BytesIO

    from docx import Document
    from docx.shared import Inches

    client, me = student
    resume_id = await _generate(client, me)
    section = Document(BytesIO((await client.get(f"{me}/built-resumes/{resume_id}/docx")).content)).sections[0]
    assert (section.page_width, section.page_height) == (Inches(8.5), Inches(11))


async def test_a_bank_fix_shows_but_a_new_entry_waits_off_the_resume(student):
    client, me = student
    await _fill_bank(client, me)
    resume_id = await _generate(client, me)
    bank = (await client.get(f"{me}/bank")).json()
    tutor = next(e for e in bank["entries"] if e["title"] == "Tutor")

    # A typo fixed in the bank shows on the resume (a live view)...
    await client.patch(f"{me}/bank/bullets/{tutor['bullets'][0]['id']}", json={"text": "Tutored 45 students in calculus"})
    # ...but an entry, bullet or skill added after Generate has no choice row.
    await client.post(f"{me}/bank/entries", json={"section": "experience", "title": "Barista", "organization": "Cafe"})
    await client.post(f"{me}/bank/entries/{tutor['id']}/bullets", json={"text": "Ran review sessions"})
    await client.post(f"{me}/bank/skills", json={"name": "Rust", "group": "Languages"})

    lines = _docx_lines((await client.get(f"{me}/built-resumes/{resume_id}/docx")).content)
    assert "Tutored 45 students in calculus" in lines
    assert "Tutored 40 students in calculus" not in lines
    assert not any("Barista" in line or "Rust" in line or "Ran review sessions" in line for line in lines)

    # A resume made now has them all.
    fresh = _docx_lines((await client.get(f"{me}/built-resumes/{await _generate(client, me)}/docx")).content)
    assert "Barista" in fresh and "Ran review sessions" in fresh and "Languages: Python, Rust" in fresh


async def test_generate_leaves_a_copy_of_an_entry_off(student):
    # An import saved "as new" left two of one job in a bank; Generate showed
    # both, and the resume ran to two pages. The copy waits in "Not on This
    # Resume"; a job held twice at one company in other months is two jobs.
    client, me = student
    job = {"section": "experience", "title": "Software Intern", "organization": "Acme",
           "start_month": "2025-06", "end_month": "2025-08", "bullets": ["Cut page load time by 30%"]}
    for body in (job, {**job, "title": "software intern ", "organization": "ACME"},
                 {**job, "start_month": "2024-06", "end_month": "2024-08", "bullets": ["Wrote tests"]}):
        res = await client.post(f"{me}/bank/entries?allow_duplicate=true", json=body)
        assert res.status_code == 201, res.text
    resume_id = await _generate(client, me)

    lines = _docx_lines((await client.get(f"{me}/built-resumes/{resume_id}/docx")).content)
    assert lines.count("Cut page load time by 30%") == 1
    assert "Wrote tests" in lines
    choices = (await client.get(f"{me}/built-resumes/{resume_id}/choices")).json()
    hidden = [e for e in choices["entries"] if not e["shown"]]
    assert len(hidden) == 1


class _FakeBucket:
    def upload(self, path, contents, options):
        return {"path": path}

    def remove(self, paths):
        return []


class _FakeSupabase:
    class storage:
        @staticmethod
        def from_(bucket):
            return _FakeBucket()


async def _upload(client, me):
    files = {"file": ("resume.pdf", b"%PDF-1.4 not really a pdf", "application/pdf")}
    return await client.post(f"{me}/resumes", files=files)


async def test_five_uploads_and_no_sixth_however_many_are_generated(student, monkeypatch):
    monkeypatch.setattr("app.routers.resumes.get_supabase", lambda: _FakeSupabase())
    monkeypatch.setattr("app.routers.resumes._extract", _no_text)
    client, me = student
    # Generated resumes count for nothing (KAN-268): six of them leave room
    # for five uploads.
    for _ in range(6):
        await _generate(client, me)
    uploads = []
    for _ in range(5):
        res = await _upload(client, me)
        assert res.status_code == 200
        uploads.append(res.json()["id"])

    res = await _upload(client, me)
    assert res.status_code == 409
    assert "5 resumes" in res.json()["detail"]
    assert (await client.post(f"{me}/built-resumes", json={})).status_code == 201

    # Deleting one makes room again.
    assert (await client.delete(f"{me}/resumes/{uploads[0]}")).status_code == 200
    assert (await _upload(client, me)).status_code == 200


async def _no_text(contents, ext):
    return ""


async def test_a_tailored_resume_says_when_its_posting_has_closed(student, database):
    client, me = student
    applicant_id = uuid.UUID(me.split("/")[2])
    base = await _generate(client, me)
    async with database() as session:
        posting = await session.scalar(
            text(
                "INSERT INTO job_postings (id, title, status, company_name, apply_url) VALUES"
                " (gen_random_uuid(), 'Intern', 'published', 'Acme', 'https://example.com/apply') RETURNING id"
            )
        )
        tailored = await session.scalar(
            text(
                "INSERT INTO built_resumes (applicant_id, name, template, job_posting_id)"
                " VALUES (:a, 'Acme – Intern', 'jakes', :p) RETURNING id"
            ),
            {"a": applicant_id, "p": posting},
        )
        await session.commit()

    async def closed() -> dict[str, bool]:
        return {r["id"]: r["posting_closed"] for r in (await client.get(f"{me}/built-resumes")).json()}

    assert await closed() == {base: False, str(tailored): False}
    async with database() as session:
        await session.execute(text("UPDATE job_postings SET status = 'closed' WHERE id = :p"), {"p": posting})
        await session.commit()
    assert await closed() == {base: False, str(tailored): True}


async def test_the_default_resume_is_an_uploaded_or_a_base_resume(student, monkeypatch):
    monkeypatch.setattr("app.routers.resumes.get_supabase", lambda: _FakeSupabase())
    monkeypatch.setattr("app.routers.resumes._extract", _no_text)
    client, me = student
    upload = (await _upload(client, me)).json()
    uploaded = upload["id"]
    assert upload["is_default"] is True  # the student's only resume
    built = await _generate(client, me)
    # A second resume leaves the default where it was.
    assert [r["is_default"] for r in (await client.get(f"{me}/built-resumes")).json()] == [False]

    assert (await client.put(f"{me}/resumes/{uploaded}/default")).status_code == 200
    assert (await client.put(f"{me}/built-resumes/{built}/default")).status_code == 200
    assert [r["is_default"] for r in (await client.get(f"{me}/built-resumes")).json()] == [True]
    assert [r["is_default"] for r in (await client.get(f"{me}/resumes")).json()] == [False]

    assert (await client.put(f"{me}/resumes/{uploaded}/default")).status_code == 200
    assert [r["is_default"] for r in (await client.get(f"{me}/built-resumes")).json()] == [False]
    assert [r["is_default"] for r in (await client.get(f"{me}/resumes")).json()] == [True]

    # Deleting the default base resume makes the one left the default.
    await client.put(f"{me}/built-resumes/{built}/default")
    await client.delete(f"{me}/built-resumes/{built}")
    assert [r["is_default"] for r in (await client.get(f"{me}/resumes")).json()] == [True]

    # And deleting an uploaded resume makes the base resume left the default.
    built = await _generate(client, me)
    await client.delete(f"{me}/resumes/{uploaded}")
    assert [r["is_default"] for r in (await client.get(f"{me}/built-resumes")).json()] == [True]


async def test_a_tailored_resume_can_be_the_default(student, database):
    """KAN-268: the API refused a tailored resume's star until then."""
    client, me = student
    base = await _generate(client, me)
    async with database() as session:
        posting = await session.scalar(
            text(
                "INSERT INTO job_postings (id, title, status, company_name, apply_url) VALUES"
                " (gen_random_uuid(), 'Intern', 'published', 'Acme', 'https://example.com/apply') RETURNING id"
            )
        )
        tailored = await session.scalar(
            text(
                "INSERT INTO built_resumes (applicant_id, name, template, job_posting_id)"
                " VALUES (:a, 'Tailored', 'jakes', :p) RETURNING id"
            ),
            {"a": uuid.UUID(me.split("/")[2]), "p": posting},
        )
        await session.commit()

    assert (await client.put(f"{me}/built-resumes/{tailored}/default")).status_code == 200
    stars = {r["id"]: r["is_default"] for r in (await client.get(f"{me}/built-resumes")).json()}
    assert stars == {str(tailored): True, base: False}
    assert (await client.put(f"{me}/built-resumes/{uuid.uuid4()}/default")).status_code == 404


async def test_an_entry_says_how_many_built_resumes_show_it(student, database):
    """Story 18: the delete confirm names the resumes an entry is on."""
    client, me = student
    await _fill_bank(client, me)
    entry = next(e for e in (await client.get(f"{me}/bank")).json()["entries"] if e["title"] == "Tutor")
    url = f"{me}/bank/entries/{entry['id']}/resumes"
    assert (await client.get(url)).json() == {"built_resumes": 0}

    first = await _generate(client, me)
    await _generate(client, me)
    assert (await client.get(url)).json() == {"built_resumes": 2}

    hide = {"entries": [{"entry_id": entry["id"], "shown": False}]}
    assert (await client.patch(f"{me}/built-resumes/{first}/choices", json=hide)).status_code == 200
    assert (await client.get(url)).json() == {"built_resumes": 1}

    async with database() as session:
        other = await _seed_applicant(session)
    assert (await client.get(f"/applicants/{other}/bank/entries/{entry['id']}/resumes")).status_code == 403
    assert (await client.get(f"{me}/bank/entries/{uuid.uuid4()}/resumes")).status_code == 404


async def test_another_students_built_resumes_are_out_of_reach(student, database):
    client, me = student
    async with database() as session:
        other = await _seed_applicant(session)
    _sign_in(other)
    theirs = await _generate(client, f"/applicants/{other}")
    _sign_in(uuid.UUID(me.split("/")[2]))

    assert (await client.get(f"/applicants/{other}/built-resumes")).status_code == 403
    assert (await client.post(f"/applicants/{other}/built-resumes", json={})).status_code == 403
    for path in ("docx", "render"):
        assert (await client.get(f"{me}/built-resumes/{theirs}/{path}")).status_code == 404
    assert (await client.put(f"{me}/built-resumes/{theirs}/default")).status_code == 404
    assert (await client.delete(f"{me}/built-resumes/{theirs}")).status_code == 404


def _tectonic_or_skip() -> None:
    from tests.test_resume_render import _tectonic_or_skip as check

    check()


def _pdf_text(encoded: str) -> str:
    import base64
    import io

    from pdfminer.high_level import extract_text

    return " ".join(extract_text(io.BytesIO(base64.b64decode(encoded))).split())


async def test_generate_then_render_gives_a_one_page_pdf_with_the_bank(student):
    _tectonic_or_skip()
    client, me = student
    await _fill_bank(client, me)
    resume_id = await _generate(client, me)

    res = await client.get(f"{me}/built-resumes/{resume_id}/render")
    assert res.status_code == 200
    body = res.json()
    assert body["page_count"] == 1
    assert body["missing_characters"] == []
    text = _pdf_text(body["pdf"])
    for expected in ("Sam Seeker", "Stony Brook University", "Software Intern", "Cut page load time by 30%",
                     # TeX's font mapping prints a straight apostrophe curly
                     "Dean’s List | Stony Brook University", "Languages: Python"):
        assert expected in text


async def test_render_gives_the_editor_a_box_for_each_line_with_bank_ids(student):
    _tectonic_or_skip()
    client, me = student
    await _fill_bank(client, me)
    resume_id = await _generate(client, me)
    bank = (await client.get(f"{me}/bank")).json()

    boxes = (await client.get(f"{me}/built-resumes/{resume_id}/render")).json()["boxes"]
    assert boxes and all(set(b) == {"item_id", "kind", "page", "x0", "y0", "x1", "y1"} for b in boxes)
    by_item = {}
    for b in boxes:
        by_item.setdefault(b["item_id"], set()).add(b["kind"])
    intern = next(e for e in bank["entries"] if e["title"] == "Software Intern")
    assert by_item[intern["id"]] == {"title", "organization", "location", "dates"}
    for bullet in intern["bullets"]:
        assert by_item[bullet["id"]] == {"bullet"}
    school = next(e for e in bank["entries"] if e["section"] == "education")
    assert by_item[school["id"]] == {"title", "location", "degree", "dates", "gpa", "coursework"}
    for skill in bank["skills"]:
        assert by_item[skill["id"]] == {"skill"}
    # PDF points from the top left of a Letter page
    assert all(b["page"] == 1 and 0 <= b["x0"] < b["x1"] <= 612 and 0 <= b["y0"] < b["y1"] <= 792 for b in boxes)


async def test_render_reports_a_long_resume_and_the_characters_the_font_lacks(student):
    _tectonic_or_skip()
    client, me = student
    bullets = [f"Organized event number {i} for two hundred members of the society" for i in range(30)]
    for title in ("Президент", "Treasurer"):
        res = await client.post(f"{me}/bank/entries", json={"section": "leadership", "title": title,
                                                            "organization": "Chess Society", "bullets": bullets})
        assert res.status_code == 201
    resume_id = await _generate(client, me)

    body = (await client.get(f"{me}/built-resumes/{resume_id}/render")).json()
    assert body["page_count"] == 2
    assert body["missing_characters"] == ["П", "р", "е", "з", "и", "д", "н", "т"]


async def test_render_answers_503_when_tectonic_is_missing(student, monkeypatch):
    from app.services.resume_render import TectonicUnavailable

    def missing(view, identity, template):
        raise TectonicUnavailable("no tectonic binary")

    monkeypatch.setattr("app.routers.built_resumes.render", missing)
    client, me = student
    resume_id = await _generate(client, me)

    res = await client.get(f"{me}/built-resumes/{resume_id}/render")
    assert res.status_code == 503
    assert res.json() == {"detail": "The PDF could not be made. Try again, or download the Word file."}


async def test_recommendations_work_when_the_default_is_a_base_resume(student, monkeypatch):
    async def no_embedding(texts, *, timeout=None):
        return None

    monkeypatch.setattr("app.queries.jobs.recommendations.embed", no_embedding)
    client, me = student
    applicant_id = uuid.UUID(me.split("/")[2])
    await _fill_bank(client, me)
    built = await _generate(client, me)
    await client.put(f"{me}/built-resumes/{built}/default")

    # The recommender reads the bank (KAN-212), whichever resume is the default.
    assert (await client.get(f"/jobs/recommended/{applicant_id}")).status_code == 200


def _at(value: str) -> datetime.datetime:
    return datetime.datetime.fromisoformat(value)


async def test_a_resume_made_with_no_name_is_numbered_past_the_highest_in_use(student):
    """KAN-258: every resume made with no name was "Base Resume"."""
    client, me = student
    first = (await client.post(f"{me}/built-resumes", json={})).json()
    second = (await client.post(f"{me}/built-resumes", json={})).json()
    assert (first["name"], second["name"]) == ("Resume #1", "Resume #2")

    # A deleted number is not handed out again while a higher one is left.
    assert (await client.delete(f"{me}/built-resumes/{first['id']}")).status_code == 204
    assert (await client.post(f"{me}/built-resumes", json={})).json()["name"] == "Resume #3"


async def test_rename_changes_the_name_and_frees_its_number(student):
    client, me = student
    resume_id = (await client.post(f"{me}/built-resumes", json={})).json()["id"]

    res = await client.patch(f"{me}/built-resumes/{resume_id}", json={"name": "  Data Science  "})
    assert res.status_code == 200
    renamed = res.json()
    assert renamed["name"] == "Data Science"
    assert _at(renamed["edited_at"]) > _at(renamed["created_at"])
    assert [r["name"] for r in (await client.get(f"{me}/built-resumes")).json()] == ["Data Science"]

    assert (await client.patch(f"{me}/built-resumes/{resume_id}", json={"name": "   "})).status_code == 422
    assert (await client.patch(f"{me}/built-resumes/{uuid.uuid4()}", json={"name": "X"})).status_code == 404
    # No resume holds "Resume #1" any more, so the next one takes it.
    assert (await client.post(f"{me}/built-resumes", json={})).json()["name"] == "Resume #1"


async def test_edited_at_is_created_at_until_the_resume_itself_changes(student):
    client, me = student
    await _fill_bank(client, me)
    resume_id = await _generate(client, me)
    made = (await client.get(f"{me}/built-resumes")).json()[0]
    assert _at(made["edited_at"]) == _at(made["created_at"])

    bank = (await client.get(f"{me}/bank")).json()
    bullet = next(e for e in bank["entries"] if e["title"] == "Software Intern")["bullets"][0]["id"]
    res = await client.patch(f"{me}/built-resumes/{resume_id}/choices",
                             json={"bullets": [{"bullet_id": bullet, "shown": False}]})
    assert res.status_code == 200

    edited = (await client.get(f"{me}/built-resumes")).json()[0]
    assert _at(edited["edited_at"]) > _at(edited["created_at"])
