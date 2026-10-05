"""DOCX text extraction: what python-docx's `paragraphs` alone loses."""
from io import BytesIO

from docx import Document

from app.routers.resumes import _extract_docx_text


def _docx_bytes(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_word_bullets_become_bullet_characters():
    # Word stores bullets as numbering properties, so `text` has no glyph
    doc = Document()
    doc.add_paragraph("Built things that scaled.", style="List Bullet")
    assert _extract_docx_text(_docx_bytes(doc)) == "• Built things that scaled."


def test_table_rows_are_read_in_document_order():
    doc = Document()
    doc.add_paragraph("EXPERIENCE")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Acme Corp"
    table.cell(0, 1).text = "Jan 2023 - Present"
    doc.add_paragraph("Backend Engineer")
    assert _extract_docx_text(_docx_bytes(doc)).splitlines() == [
        "EXPERIENCE",
        "Acme Corp  Jan 2023 - Present",
        "Backend Engineer",
    ]


def test_tabs_become_two_space_column_gaps():
    # The parser reads a 2+ space gap as a column break (e.g. before a location)
    doc = Document()
    doc.add_paragraph("Software Engineer\t\tBrooklyn, NY")
    assert _extract_docx_text(_docx_bytes(doc)) == "Software Engineer  Brooklyn, NY"
